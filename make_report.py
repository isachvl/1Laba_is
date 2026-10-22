"""Generate evidence CSVs, two Optuna plots, and a three-page Word report after the run."""
import csv
import json
import statistics as stats
from collections import Counter
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import mlflow
import optuna
import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path("report")
ROOT.mkdir(exist_ok=True)


def get_run_entries():
    rows = []
    for metadata_file in Path("runs").glob("*/metadata.json"):
        metrics_file = metadata_file.parent / "metrics.json"
        if not metrics_file.exists():
            continue
        try:
            m = json.loads(metadata_file.read_text(encoding="utf-8"))
            d = json.loads(metrics_file.read_text(encoding="utf-8"))
            if m.get("status") not in ("completed", "complete"):
                continue
            rows.append({**m, **d, "path": metadata_file.parent,
                         "mtime": metadata_file.stat().st_mtime})
        except (ValueError, OSError):
            pass
    return rows


def select(runs, seed, series=None, variant=None):
    candidates = [r for r in runs if r.get("seed") == seed and
                  (series is None or r.get("series") == series) and
                  (variant is None or r.get("variant") == variant)]
    return max(candidates, key=lambda r: r["mtime"]) if candidates else None


def prediction_map(run, split="test"):
    if run is None:
        return {}
    with (run["path"] / "predictions.csv").open(encoding="utf-8", newline="") as f:
        return {int(r["index"]): (int(r["y_true"]), int(r["y_pred"]))
                for r in csv.DictReader(f) if r["split"] == split}


def differences(a, b, split):
    p, q = prediction_map(a, split), prediction_map(b, split)
    if p.keys() != q.keys() or not p:
        raise ValueError(f"Mismatch in {split} image indices")
    corrected = worsened = 0
    for i in p:
        assert p[i][0] == q[i][0]
        old_ok = p[i][0] == p[i][1]
        new_ok = q[i][0] == q[i][1]
        corrected += (not old_ok and new_ok)
        worsened += (old_ok and not new_ok)
    return corrected, worsened


def baseline_summary(runs):
    # Previous baseline_seed1..5 records from the user's existing runs are supported.
    selected = {}
    for seed in range(1, 6):
        matching = [r for r in runs if r.get("seed") == seed and
                    r.get("series") in ("base_seeds", "baseline") and
                    r.get("variant") in (None, "", "base")]
        if matching:
            selected[seed] = max(matching, key=lambda r: r["mtime"])
    values = [r["val_acc"] for r in selected.values()]
    if values:
        spread = max(values) - min(values)
        description = (f"min {min(values)*100:.2f}%, median {stats.median(values)*100:.2f}%, "
                       f"max {max(values)*100:.2f}%; range {spread*100:.2f} pp")
    else:
        spread, description = None, "no seed 1..5 runs found"
    return selected, spread, description


def collect_optuna():
    if not Path("optuna.db").exists():
        return None
    study = optuna.load_study(study_name="lr0_search", storage="sqlite:///optuna.db")
    study.trials_dataframe().to_csv(ROOT / "trials.csv", index=False)
    counts = Counter(t.state.name for t in study.trials)
    epochs_run = 0
    time_run = 0.0
    complete_epoch_times = []
    for metadata_file in Path("runs").glob("*/metadata.json"):
        try:
            m = json.loads(metadata_file.read_text(encoding="utf-8"))
            if m.get("series") != "search":
                continue
            history = metadata_file.parent / "history.csv"
            if not history.exists():
                continue
            with history.open(newline="", encoding="utf-8") as f:
                history_rows = list(csv.DictReader(f))
            epochs_run += len(history_rows)
            times = [float(r["epoch_time"]) for r in history_rows]
            time_run += sum(times)
            if m.get("status") in ("completed", "complete"):
                complete_epoch_times.extend(times)
        except (ValueError, OSError):
            continue
    all_budget = 30 * len(study.trials)
    saved_epochs = max(0, all_budget - epochs_run)
    estimated_saved_s = (stats.mean(complete_epoch_times) * saved_epochs
                         if complete_epoch_times else 0.0)
    # Plot 1: completed trial val accuracy + best-so-far curve.
    nums = []
    values = []
    for t in study.trials:
        if t.state.name == "COMPLETE" and t.value is not None:
            nums.append(t.number)
            values.append(float(t.value) * 100)
    if nums:
        bests = []
        vbest = float("-inf")
        for v in values:
            vbest = max(vbest, v)
            bests.append(vbest)
        plt.figure(figsize=(8, 3.4))
        plt.scatter(nums, values, s=20, label="Completed trial")
        plt.step(nums, bests, where="post", label="Best so far")
        plt.xlabel("Trial number")
        plt.ylabel("Validation accuracy, %")
        plt.legend()
        plt.tight_layout()
        plt.savefig(ROOT / "optuna_accuracy.png", dpi=160)
        plt.close()
    try:
        importance = optuna.importance.get_param_importances(study)
    except (ValueError, RuntimeError):
        importance = {}
    if importance:
        plt.figure(figsize=(8, 3.1))
        plt.barh(list(importance.keys())[::-1], list(importance.values())[::-1])
        plt.xlabel("Importance")
        plt.tight_layout()
        plt.savefig(ROOT / "optuna_importance.png", dpi=160)
        plt.close()
    return dict(study=study, counts=counts, epochs_run=epochs_run,
                saved_epochs=saved_epochs, measured_time=time_run,
                saved_time=estimated_saved_s, importance=importance)


def collect_final(runs):
    pairs = []
    for seed in range(101, 106):
        base = select(runs, seed, series="final", variant="base")
        best = select(runs, seed, series="final", variant="best")
        if base is not None and best is not None:
            cv, wv = differences(base, best, "val")
            ct, wt = differences(base, best, "test")
            pairs.append(dict(seed=seed, base=base, best=best,
                              delta_val=best["val_acc"]-base["val_acc"],
                              delta_test=best["test_acc"]-base["test_acc"],
                              fixed_val=cv, broken_val=wv,
                              fixed_test=ct, broken_test=wt))
    return pairs


def write_csv(rows, path, fields):
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(fields)
        for row in rows:
            writer.writerow(row)


def get_hypothesis():
    """Берём прогноз из твоего hypothesis.md, а не придумываем задним числом."""
    path = Path("hypothesis.md")
    if not path.exists() or not path.read_text(encoding="utf-8").strip():
        return "Прогноз не записан. Заполни hypothesis.md до запуска поиска Optuna."
    lines = path.read_text(encoding="utf-8").splitlines()
    # Markdown-заголовки в отчёте превращаем в обычный текст.
    return " ".join(line.lstrip("#- ").strip() for line in lines if line.strip())[:1600]


def create_docx(baseline, base_desc, opt, pairs, spread, md_lines):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    def heading(value, level=1):
        doc.add_heading(value, level=level)
    def para(value):
        doc.add_paragraph(value)
    def table(headers, data):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Shading Accent 1"
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = str(h)
        for row in data:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
    doc.add_heading("ЛР0. Воспроизводимый эксперимент", 0)
    para("Задача: проверить, даёт ли подбор гиперпараметров улучшение больше случайного разброса.")
    heading("Прогноз и постановка")
    para("Зафиксированный прогноз: " + get_hypothesis())
    para("Данные: sklearn Digits; 1077 train / 360 val / 360 test. Разделение фиксировано, "
         "масштаб признаков — деление на 16. Модель — MLP, 30 эпох, CPU, Adam.")
    heading("Эксперимент 1. Разброс по seed")
    table(["Seed", "val_acc", "val_loss", "Время, с"], [
        (s, f"{r['val_acc']*100:.2f}%", f"{r['val_loss']:.4f}", f"{r['train_time']:.2f}")
        for s, r in sorted(baseline.items())
    ])
    para("Сводка: " + base_desc + ".")
    if 1 in baseline and 2 in baseline:
        a, b = differences(baseline[1], baseline[2], "val")
        para(f"Сравнение ошибок seed 1 и 2: второй исправил {a} предсказаний первого "
             f"и испортил {b} правильных предсказаний первого.")
    heading("Повторяемость")
    para("Одинаковый seed даёт совпадающие предсказания и значения train_loss/val_loss/val_acc "
         "при одинаковых данных, коде и окружении. Поле epoch_time в CSV не обязано совпадать "
         "побайтно, поскольку фиксирует реальное время исполнения.")
    doc.add_page_break()
    heading("Эксперимент 2. Поиск Optuna")
    if opt is None:
        para("Поиск ещё не выполнен.")
    else:
        study = opt["study"]
        c = opt["counts"]
        para(f"Study lr0_search: {len(study.trials)} trial; COMPLETE={c['COMPLETE']}, "
             f"PRUNED={c['PRUNED']}, FAIL={c['FAIL']}. "
             "TPE(seed=42), MedianPruner(5, 5), 30 эпох, seed обучения 0.")
        para(f"Выполнено эпох: {opt['epochs_run']} из {30 * len(study.trials)}; "
             f"сэкономлено по бюджету {opt['saved_epochs']} эпох (включая досрочно остановленные и "
             f"не завершившиеся trial). Оценка времени сохранённых эпох: {opt['saved_time']:.1f} с "
             "по средней скорости завершённых trial.")
        if any(t.state.name == "COMPLETE" for t in study.trials):
            para(f"Лучший trial #{study.best_trial.number}: "
                 f"val_acc={study.best_value*100:.2f}%; "
                 f"параметры {json.dumps(study.best_params, ensure_ascii=False)}")
    for image in ("optuna_accuracy.png", "optuna_importance.png"):
        if (ROOT / image).exists():
            doc.add_picture(str(ROOT / image), width=Cm(14))
    doc.add_page_break()
    heading("Эксперимент 3. Честное сравнение")
    if pairs:
        table(["Seed", "Δ val, п.п.", "Δ test, п.п.", "Исправлено test", "Испорчено test"], [
            (p["seed"], f"{100*p['delta_val']:+.2f}", f"{100*p['delta_test']:+.2f}",
             p["fixed_test"], p["broken_test"]) for p in pairs
        ])
        v = [p["delta_val"] for p in pairs]
        t = [p["delta_test"] for p in pairs]
        para(f"Δ val: медиана {100*stats.median(v):+.2f} п.п., диапазон "
             f"[{100*min(v):+.2f}; {100*max(v):+.2f}] п.п. "
             f"Δ test: медиана {100*stats.median(t):+.2f} п.п., диапазон "
             f"[{100*min(t):+.2f}; {100*max(t):+.2f}] п.п.")
        para(f"Всего исправлено на test: {sum(p['fixed_test'] for p in pairs)}; "
             f"испорчено: {sum(p['broken_test'] for p in pairs)}.")
        if opt and any(x.state.name == "COMPLETE" for x in opt["study"].trials):
            para(f"Лучший trial в поиске val={100*opt['study'].best_value:.2f}%; "
                 f"новые seed лучшей модели: val={100*stats.mean(p['best']['val_acc'] for p in pairs):.2f}%, "
                 f"test={100*stats.mean(p['best']['test_acc'] for p in pairs):.2f}%. "
                 "Снижение после поиска возможно из-за смещения отбора по val.")
        if len(pairs) == 5 and spread is not None and all(x > spread for x in t):
            conclusion = "Все пять положительных улучшений test больше разброса baseline: эффект устойчив в рамках проведённой проверки."
        elif any(x < 0 for x in t) and any(x > 0 for x in t):
            conclusion = "Разности на test имеют разные знаки: данных недостаточно для уверенного утверждения об улучшении."
        else:
            conclusion = "Эффект не соответствует строгому критерию: пять положительных разностей на test больше разброса baseline."
        para("Вывод: " + conclusion)
    else:
        para("Финальные запуски с seed 101–105 ещё не выполнены. После них отчёт заполнится автоматически.")
    heading("Воспроизводимость и помощь ИИ")
    para("Все конфигурации, split SHA256, git commit и версии пакетов сохранены; точность "
         "пересчитывается из predictions.csv, история — из history.csv. "
         "Код каркаса и отчёта готовился с помощью ChatGPT; проверка — запуск, "
         "assert-проверки, сравнение сохранённых CSV и графиков MLflow. "
         "Итоговые численные выводы следует сверить с выгрузкой реальных запусков.")
    out = ROOT / "LR0_report.docx"
    doc.save(out)
    print("Created", out)


def main():
    runs = get_run_entries()
    baseline, spread, desc = baseline_summary(runs)
    opt = collect_optuna()
    pairs = collect_final(runs)
    write_csv([(s, r["val_acc"], r["val_loss"], r["train_time"])
               for s, r in sorted(baseline.items())],
              ROOT / "baseline.csv", ["seed", "val_acc", "val_loss", "train_time"])
    write_csv([(p["seed"], p["delta_val"], p["delta_test"], p["fixed_test"], p["broken_test"])
               for p in pairs], ROOT / "final_deltas.csv",
              ["seed", "delta_val", "delta_test", "corrected_test", "worsened_test"])
    try:
        mlflow.set_tracking_uri("sqlite:///mlflow.db")
        ex = mlflow.get_experiment_by_name("lr0")
        if ex:
            mlflow.search_runs(experiment_ids=[ex.experiment_id]).to_csv(ROOT / "mlflow_export.csv", index=False)
    except Exception as e:
        print("MLflow CSV export skipped:", e)
    md = ["# ЛР0 — отчёт", "", "## Прогноз",
          get_hypothesis(),
          "", "## Эксперимент 1", "Разброс: " + desc, "",
          "| seed | val_acc | val_loss | время (с) |", "|---|---:|---:|---:|"]
    for s, r in sorted(baseline.items()):
        md.append(f"|{s}|{r['val_acc']:.4f}|{r['val_loss']:.4f}|{r['train_time']:.2f}|")
    md.extend(["", "## Эксперимент 2 (Optuna)"])
    if opt:
        md.extend([f"Trials: {len(opt['study'].trials)}, состояния: {dict(opt['counts'])}.",
                   f"Эпох: {opt['epochs_run']}, сэкономлено эпох: {opt['saved_epochs']}, "
                   f"оценка сэкономленного времени: {opt['saved_time']:.1f} с."])
        if any(t.state.name == "COMPLETE" for t in opt["study"].trials):
            md.append(f"Лучший trial: #{opt['study'].best_trial.number}, "
                      f"val_acc={opt['study'].best_value:.4f}.")
        md.extend(["![Trial accuracy](optuna_accuracy.png)", "![Importance](optuna_importance.png)"])
    else:
        md.append("Пока нет данных поиска.")
    md.extend(["", "## Эксперимент 3 (финальные seed 101–105)",
               "| seed | Δval (п.п.) | Δtest (п.п.) | test исправлено | test испорчено |",
               "|---|---:|---:|---:|---:|"])
    for p in pairs:
        md.append(f"|{p['seed']}|{100*p['delta_val']:+.2f}|{100*p['delta_test']:+.2f}|"
                  f"{p['fixed_test']}|{p['broken_test']}|")
    if len(pairs) == 5:
        dv = [p['delta_val'] for p in pairs]
        dt = [p['delta_test'] for p in pairs]
        md.append(f"Медиана Δval={100*stats.median(dv):+.2f} п.п. (диапазон {100*min(dv):+.2f}...{100*max(dv):+.2f}); "
                  f"медиана Δtest={100*stats.median(dt):+.2f} п.п. (диапазон {100*min(dt):+.2f}...{100*max(dt):+.2f}).")
        md.append(f"На test исправлено {sum(p['fixed_test'] for p in pairs)}, "
                  f"испорчено {sum(p['broken_test'] for p in pairs)} по всем пяти seed.")
    else:
        md.append("Финальные результаты ещё не полные: запускать make_report.py повторно после run_final.py.")
    md.extend(["", "## Помощь ИИ и проверка", "Код создавался при помощи ChatGPT; метрики "
               "пересчитаны из сохранённых CSV, тесты реализации и реальные запуски выполнены самостоятельно."])
    (ROOT / "generated_report.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    create_docx(baseline, desc, opt, pairs, spread, md)
    print("Report files in report/. final paired rows:", len(pairs))


if __name__ == "__main__":
    main()
